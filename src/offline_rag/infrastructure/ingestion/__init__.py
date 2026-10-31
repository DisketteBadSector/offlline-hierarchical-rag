"""Document Ingestion Adapters (PDF, DOCX, DOC) using Docling and LVM Layout Assistant (Section 7.3)."""

from __future__ import annotations

from collections.abc import Mapping, Sequence
from functools import lru_cache
import hashlib
from io import BytesIO
from pathlib import Path
import re
from typing import Any, Protocol
import zipfile

import docx
from docling.datamodel.base_models import DocumentStream, InputFormat
from docling.datamodel.pipeline_options import (
    NativePdfPipelineOptions,
    PdfPipelineOptions,
)
from docling.document_converter import (
    DocumentConverter,
    NativePdfFormatOption,
    PdfFormatOption,
    WordFormatOption,
)
from docling_core.types.doc.labels import DocItemLabel

from offline_rag.application.ports import (
    DocumentExtractorPort,
    LvmLayoutAssistantPort,
)
from offline_rag.domain.knowledge import (
    DocumentFormat,
    ExtractedBlock,
    ExtractedDocumentContent,
    ExtractedTable,
)
from offline_rag.domain.shared.errors import DomainValidationError
from offline_rag.domain.shared.value_objects import normalize_persian_text

OLE2_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
NUMBERED_HEADING_RE = re.compile(r"^([۰-۹0-9]+[\.\-\)]+)\s+.+$")
DOC_STYLE_NOISE = frozenset(
    {
        "normal",
        "standard",
        "default style",
        "default paragraph font",
        "table normal",
        "table contents",
        "table heading",
        "no list",
        "list",
        "numbering symbols",
        "horizontal line",
        "preformatted text",
        "body text",
        "heading",
        "heading 1",
        "heading 2",
        "heading 3",
        "heading 4",
        "title",
        "subtitle",
        "caption",
        "index",
        "header",
        "footer",
        "footnote",
        "endnote",
        "quotations",
        "times new roman",
        "arial",
        "calibri",
        "tahoma",
        "b nazanin",
        "vazirmatn",
        "liberationsans",
        "liberation serif",
        "liberation sans",
        "liberation mono",
        "freesans",
        "dejavu sans",
        "microsoft word",
        "root entry",
        "worddocument",
        "summaryinformation",
        "documentsummaryinformation",
        "compobj",
        "1table",
        "0table",
    }
)


class DoclingConverterProtocol(Protocol):
    """Minimal protocol for Docling's DocumentConverter."""

    def convert(self, source: Any, *, raises_on_error: bool = True) -> Any:
        """Convert a file path or DocumentStream into a ConversionResult."""
        ...


def _sha256_hex(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _to_int(value: object, default: int = 1) -> int:
    """Safely convert an arbitrary object to int for strict static type checkers."""
    if isinstance(value, bool):
        return int(value)
    if isinstance(value, int):
        return value
    if isinstance(value, (str, bytes, bytearray)):
        try:
            return int(value)
        except ValueError:
            return default
    return default


def _update_heading_stack(
    stack: list[tuple[int, str]], level: int, title: str
) -> tuple[str, ...]:
    while stack and stack[-1][0] >= level:
        stack.pop()
    stack.append((level, title.strip()))
    return tuple(item[1] for item in stack)


@lru_cache(maxsize=4)
def build_offline_docling_converter(
    artifacts_path: str | None = None,
) -> DocumentConverter:
    """Create and cache a Docling DocumentConverter configured for 100% offline RAG execution.

    - Uses ``NativePdfFormatOption`` (``ThreadedDoclingParseDocumentBackend`` + ``NativePdfPipeline``)
      by default so PDF extraction runs completely offline with zero HuggingFace network calls.
    - When ``artifacts_path`` is provided with local model weights, switches to ``PdfFormatOption``
      with ``do_ocr=False`` and ``do_table_structure=True``.
    - Disables remote services and page image rendering to respect the 6GB VRAM budget.
    """
    if artifacts_path:
        pdf_options = PdfPipelineOptions(
            artifacts_path=Path(artifacts_path),
            enable_remote_services=False,
            do_ocr=False,
            do_table_structure=True,
            generate_page_images=False,
            generate_picture_images=False,
        )
        pdf_format_option: PdfFormatOption | NativePdfFormatOption = PdfFormatOption(
            pipeline_options=pdf_options
        )
    else:
        native_options = NativePdfPipelineOptions(
            enable_remote_services=False,
            generate_page_images=False,
            generate_picture_images=False,
        )
        pdf_format_option = NativePdfFormatOption(pipeline_options=native_options)

    return DocumentConverter(
        allowed_formats=[InputFormat.PDF, InputFormat.DOCX],
        format_options={
            InputFormat.PDF: pdf_format_option,
            InputFormat.DOCX: WordFormatOption(),
        },
    )


def _extract_page_no(item: object, default_page: int = 1) -> int:
    prov = getattr(item, "prov", None)
    if isinstance(prov, Sequence) and len(prov) > 0:
        first_prov: object = prov[0]
        page_no = getattr(first_prov, "page_no", None)
        if isinstance(page_no, int) and page_no >= 1:
            return page_no
    return default_page


def _extract_table_rows(item: object) -> tuple[tuple[str, ...], ...]:
    data = getattr(item, "data", None)
    grid = getattr(data, "grid", None)
    if isinstance(grid, Sequence):
        rows: list[tuple[str, ...]] = []
        for row in grid:
            if not isinstance(row, Sequence):
                continue
            cells = tuple(
                normalize_persian_text(str(getattr(cell, "text", cell)))
                for cell in row
            )
            if any(cells):
                rows.append(cells)
        return tuple(rows)
    return ()


def _legacy_doc_to_docx_bytes(raw_bytes: bytes, document_id: str) -> bytes:
    """Convert legacy binary OLE2 or plain-text .doc bytes into an in-memory .docx stream for Docling."""
    if zipfile.is_zipfile(BytesIO(raw_bytes)):
        return raw_bytes

    paragraphs: list[str] = []
    if raw_bytes.startswith(OLE2_MAGIC):
        text16 = raw_bytes.decode("utf-16le", errors="ignore")
        runs: list[str] = re.findall(
            r"[\u0600-\u06FFa-zA-Z0-9\s\.,:;\-\(\)\[\]/«»٪\u200c\r\n—–_@#]{8,}",
            text16,
        )
        for run in runs:
            for raw_para in re.split(r"[\r\n]+", run):
                cleaned = normalize_persian_text(raw_para)
                low = cleaned.lower()
                if (
                    len(cleaned) < 4
                    or low in DOC_STYLE_NOISE
                    or any(
                        low.startswith(p)
                        for p in ("normal ", "heading ", "toc ", "hyperlink")
                    )
                ):
                    continue
                if (
                    re.search(r"[\u0600-\u06FF]{2,}", cleaned)
                    or len(re.findall(r"[A-Za-z]{3,}", cleaned)) >= 2
                ):
                    if cleaned not in paragraphs:
                        paragraphs.append(cleaned)
    else:
        for enc in ("utf-8", "utf-16", "cp1256"):
            try:
                decoded = raw_bytes.decode(enc)
                paragraphs = [
                    normalize_persian_text(p)
                    for p in re.split(r"[\r\n]+", decoded)
                    if normalize_persian_text(p)
                ]
                if paragraphs:
                    break
            except Exception:
                continue

    if not paragraphs:
        raise DomainValidationError(
            f"Could not extract readable text from DOC file '{document_id}'."
        )

    word_doc = docx.Document()
    for idx, para in enumerate(paragraphs):
        if idx == 0 and len(para) <= 85:
            word_doc.add_heading(para, level=1)
        elif (
            len(para) <= 55
            and not para.endswith(".")
            and NUMBERED_HEADING_RE.match(para)
        ):
            word_doc.add_heading(para, level=2)
        else:
            word_doc.add_paragraph(para)

    buf = BytesIO()
    word_doc.save(buf)
    return buf.getvalue()


def _map_docling_document(
    doc: Any,
    document_id: str,
    file_format: DocumentFormat,
    content_hash: str,
    min_chars_per_page_before_lvm: int = 40,
) -> ExtractedDocumentContent:
    """Map a DoclingDocument into the domain's ExtractedDocumentContent value object."""
    heading_stack: list[tuple[int, str]] = []
    current_headings: tuple[str, ...] = ()
    blocks: list[ExtractedBlock] = []
    tables: list[ExtractedTable] = []
    flagged_pages: set[int] = set()
    chars_per_page: dict[int, int] = {}
    short_fragments_per_page: dict[int, int] = {}
    tables_per_page: dict[int, int] = {}
    max_seen_page = 1
    seen_first_on_page1 = False

    for entry in doc.iterate_items():
        if isinstance(entry, tuple) and len(entry) >= 2:
            item: object = entry[0]
            tree_level = _to_int(entry[1], default=1)
        else:
            item = entry
            tree_level = 1

        raw_label = getattr(item, "label", DocItemLabel.TEXT)
        label = str(getattr(raw_label, "value", raw_label)).lower()
        page_no = _extract_page_no(item, default_page=1)
        max_seen_page = max(max_seen_page, page_no)

        if label in (
            DocItemLabel.PICTURE.value,
            DocItemLabel.CHART.value,
            DocItemLabel.FORMULA.value,
            "figure",
        ):
            flagged_pages.add(page_no)
            continue

        if label == DocItemLabel.TABLE.value:
            rows = _extract_table_rows(item)
            if rows:
                caption = current_headings[-1] if current_headings else None
                tbl = ExtractedTable(
                    rows=rows,
                    page_number=page_no,
                    heading_path=current_headings,
                    caption=caption,
                )
                tables.append(tbl)
                tables_per_page[page_no] = tables_per_page.get(page_no, 0) + 1
                tbl_text = tbl.to_text()
                chars_per_page[page_no] = chars_per_page.get(page_no, 0) + len(tbl_text)
                blocks.append(
                    ExtractedBlock(
                        text=tbl_text,
                        heading_path=current_headings,
                        page_number=page_no,
                        section_reference=caption or f"page-{page_no}",
                        block_type="table",
                    )
                )
            continue

        raw_text_attr = getattr(item, "text", None)
        raw_text = normalize_persian_text(
            str(raw_text_attr) if raw_text_attr is not None else ""
        )
        if not raw_text:
            continue
        chars_per_page[page_no] = chars_per_page.get(page_no, 0) + len(raw_text)
        if len(raw_text) <= 22:
            short_fragments_per_page[page_no] = (
                short_fragments_per_page.get(page_no, 0) + 1
            )

        is_explicit_heading = label in (
            DocItemLabel.TITLE.value,
            DocItemLabel.SECTION_HEADER.value,
            "heading",
        )
        is_inferred_pdf_heading = (
            file_format == DocumentFormat.PDF
            and not is_explicit_heading
            and (
                (page_no == 1 and not seen_first_on_page1 and len(raw_text) <= 90)
                or (
                    len(raw_text) <= 75
                    and not raw_text.endswith(".")
                    and NUMBERED_HEADING_RE.match(raw_text) is not None
                )
            )
        )
        if page_no == 1:
            seen_first_on_page1 = True

        if is_explicit_heading or is_inferred_pdf_heading:
            if label == DocItemLabel.TITLE.value or (
                is_inferred_pdf_heading and not current_headings
            ):
                lvl = 1
            elif is_inferred_pdf_heading:
                lvl = 2
            else:
                raw_level = getattr(item, "level", None)
                parsed_level = _to_int(
                    raw_level if raw_level is not None else tree_level,
                    default=tree_level,
                )
                lvl = max(1, min(parsed_level, 6))
            current_headings = _update_heading_stack(heading_stack, lvl, raw_text)
            blocks.append(
                ExtractedBlock(
                    text=raw_text,
                    heading_path=current_headings,
                    page_number=page_no,
                    section_reference=current_headings[-1],
                    block_type="heading",
                    heading_level=lvl,
                )
            )
        else:
            if (
                file_format == DocumentFormat.PDF
                and blocks
                and blocks[-1].block_type == "paragraph"
                and blocks[-1].page_number == page_no
                and blocks[-1].heading_path == current_headings
            ):
                prev = blocks[-1]
                blocks[-1] = ExtractedBlock(
                    text=f"{prev.text} {raw_text}",
                    heading_path=prev.heading_path,
                    page_number=prev.page_number,
                    section_reference=prev.section_reference,
                    block_type="paragraph",
                )
            else:
                blocks.append(
                    ExtractedBlock(
                        text=raw_text,
                        heading_path=current_headings,
                        page_number=page_no,
                        section_reference=(
                            current_headings[-1]
                            if current_headings
                            else (
                                f"page-{page_no}"
                                if file_format == DocumentFormat.PDF
                                else None
                            )
                        ),
                        block_type="paragraph",
                    )
                )

    num_pages_fn = getattr(doc, "num_pages", None)
    raw_pages: object = num_pages_fn() if callable(num_pages_fn) else 1
    total_pages = max(_to_int(raw_pages, default=1), max_seen_page, 1)

    if file_format == DocumentFormat.PDF:
        for p in range(1, total_pages + 1):
            if chars_per_page.get(p, 0) < min_chars_per_page_before_lvm:
                flagged_pages.add(p)
            elif (
                short_fragments_per_page.get(p, 0) >= 35
                and tables_per_page.get(p, 0) == 0
            ):
                flagged_pages.add(p)

    sorted_flagged = tuple(sorted(flagged_pages))
    return ExtractedDocumentContent(
        document_id=document_id,
        file_format=file_format,
        blocks=tuple(blocks),
        tables=tuple(tables),
        page_count=total_pages,
        requires_lvm_assistance=len(sorted_flagged) > 0,
        flagged_lvm_pages=sorted_flagged,
        raw_content_hash=content_hash,
    )


class _DoclingExtractorBase(DocumentExtractorPort):
    """Shared Docling-backed extractor implementation for PDF and Word documents."""

    def __init__(
        self,
        converter: DoclingConverterProtocol | None = None,
        artifacts_path: str | None = None,
        min_chars_per_page_before_lvm: int = 40,
    ) -> None:
        self._converter = converter
        self._artifacts_path = artifacts_path
        self._min_chars = min_chars_per_page_before_lvm

    def _get_converter(self) -> DoclingConverterProtocol:
        if self._converter is None:
            self._converter = build_offline_docling_converter(self._artifacts_path)
        return self._converter

    def extract(
        self, file_path: Path, document_id: str | None = None
    ) -> ExtractedDocumentContent:
        if not file_path.exists():
            raise FileNotFoundError(
                f"{self.supported_format.value} file not found: {file_path}"
            )
        doc_id = document_id or file_path.stem
        return self.extract_from_bytes(file_path.read_bytes(), document_id=doc_id)

    def extract_from_bytes(
        self, raw_bytes: bytes, document_id: str
    ) -> ExtractedDocumentContent:
        if not raw_bytes:
            raise DomainValidationError(f"Document '{document_id}' is empty.")

        content_hash = _sha256_hex(raw_bytes)
        if self.supported_format == DocumentFormat.DOC:
            stream_bytes = _legacy_doc_to_docx_bytes(raw_bytes, document_id)
            ext = "docx"
        elif self.supported_format == DocumentFormat.DOCX:
            stream_bytes = raw_bytes
            ext = "docx"
        else:
            stream_bytes = raw_bytes
            ext = "pdf"

        stream = DocumentStream(
            name=f"{document_id}.{ext}", stream=BytesIO(stream_bytes)
        )
        try:
            result = self._get_converter().convert(stream, raises_on_error=True)
        except DomainValidationError:
            raise
        except Exception as exc:
            raise DomainValidationError(
                f"Docling failed to extract '{document_id}': {exc}"
            ) from exc

        return _map_docling_document(
            doc=result.document,
            document_id=document_id,
            file_format=self.supported_format,
            content_hash=content_hash,
            min_chars_per_page_before_lvm=self._min_chars,
        )


class DocxExtractorAdapter(_DoclingExtractorBase):
    """Docling-based adapter extracting paragraphs, heading hierarchies, and tables from .docx files."""

    @property
    def supported_format(self) -> DocumentFormat:
        return DocumentFormat.DOCX


class PdfExtractorAdapter(_DoclingExtractorBase):
    """Docling-based adapter extracting page-indexed blocks and tables from .pdf files with LVM flagging."""

    @property
    def supported_format(self) -> DocumentFormat:
        return DocumentFormat.PDF


class DocExtractorAdapter(_DoclingExtractorBase):
    """Docling-based adapter for .doc Word documents."""

    @property
    def supported_format(self) -> DocumentFormat:
        return DocumentFormat.DOC


class MockLvmLayoutAdapter(LvmLayoutAssistantPort):
    """Offline Mock/Fallback Adapter for LVM visual and layout interpretation (Section 7.3)."""

    def __init__(self) -> None:
        self.processed_requests: list[dict[str, Any]] = []

    def describe_visual_page(
        self, document_id: str, version_id: str, page_number: int
    ) -> str:
        self.processed_requests.append(
            {
                "document_id": document_id,
                "version_id": version_id,
                "page_number": page_number,
            }
        )
        return (
            f"[تفسیر بصری LVM - سند {document_id} نسخه {version_id} صفحه {page_number}]: "
            f"چیدمان صفحه شامل جدول مقایسه‌ای شاخص‌ها و نمودار ساختاری است که به متن نسخه متصل شد."
        )

    def enrich_extracted_content(
        self, extracted: ExtractedDocumentContent, version_id: str = "v1"
    ) -> ExtractedDocumentContent:
        if not extracted.requires_lvm_assistance:
            return extracted

        target_pages = extracted.flagged_lvm_pages or (1,)
        lvm_blocks: list[ExtractedBlock] = []
        lvm_notes: list[str] = []

        for page_num in target_pages:
            desc = self.describe_visual_page(
                document_id=extracted.document_id,
                version_id=version_id,
                page_number=page_num,
            )
            lvm_notes.append(desc)
            lvm_blocks.append(
                ExtractedBlock(
                    text=desc,
                    heading_path=("تفسیر چیدمان بصری LVM",),
                    page_number=page_num,
                    section_reference=f"lvm-page-{page_num}",
                    block_type="lvm_layout",
                )
            )

        return extracted.with_lvm_enrichment(
            lvm_blocks=lvm_blocks, lvm_notes=lvm_notes
        )


class DocumentExtractorRegistry:
    """Strategy registry selecting the appropriate DocumentExtractorPort by file extension or format."""

    def __init__(
        self,
        extractors: Mapping[DocumentFormat, DocumentExtractorPort] | None = None,
        converter: DoclingConverterProtocol | None = None,
    ) -> None:
        default_map: dict[DocumentFormat, DocumentExtractorPort] = {
            DocumentFormat.DOC: DocExtractorAdapter(converter=converter),
            DocumentFormat.DOCX: DocxExtractorAdapter(converter=converter),
            DocumentFormat.PDF: PdfExtractorAdapter(converter=converter),
        }
        if extractors:
            default_map.update(extractors)
        self._extractors = default_map

    def get_extractor(
        self, format_or_suffix: DocumentFormat | str
    ) -> DocumentExtractorPort:
        if isinstance(format_or_suffix, DocumentFormat):
            fmt = format_or_suffix
        else:
            cleaned = str(format_or_suffix).strip().upper().lstrip(".")
            try:
                fmt = DocumentFormat(cleaned)
            except ValueError as exc:
                raise DomainValidationError(
                    f"Unsupported document format or extension: '{format_or_suffix}'."
                ) from exc

        extractor = self._extractors.get(fmt)
        if extractor is None:
            raise DomainValidationError(
                f"No DocumentExtractorPort registered for {fmt.value}."
            )
        return extractor


__all__ = [
    "DocExtractorAdapter",
    "DoclingConverterProtocol",
    "DocumentExtractorPort",
    "DocumentExtractorRegistry",
    "DocxExtractorAdapter",
    "LvmLayoutAssistantPort",
    "MockLvmLayoutAdapter",
    "PdfExtractorAdapter",
    "build_offline_docling_converter",
]